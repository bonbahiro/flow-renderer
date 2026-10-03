# -*- coding: utf-8 -*-
"""
flow_renderer.py  —  業務フロー図（Layer-1〜3）を JSON 仕様から編集可能な PowerPoint に描画する

使い方:
    from flow_renderer import render
    warnings = render(spec_dict, "out.pptx")      # spec は dict または JSON ファイルパス
    for w in warnings: print(w)

特徴:
- すべて PowerPoint のネイティブ図形。矢印は「カギ線コネクタ」で図形に接続済みなので、
  PowerPoint 上で図形を動かすと矢印が自動で追従する。
- 座標はグリッド単位（既定 48列×30行、A3横）。モデルはピクセル計算をしなくてよい。
- 描画後に検査（ID不整合・重なり・領域はみ出し・孤立ノード・矢印がノードを横切る）を行い、
  警告リストを返す。警告がゼロになるまで座標を直すのが推奨運用。

依存: python-pptx (>=0.6.21)
"""
import json
from copy import deepcopy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn
from lxml import etree

# ------------------------------------------------------------------ スタイル定義（共通・変更禁止）
FONT = "Meiryo UI"

SYSTEM_COLORS = {            # システム区分 → 塗り色
    "sales":     "FEF9C3",   # 販売・案件管理
    "bi":        "DCFCE7",   # 分析・BI
    "erp":       "E0F2FE",   # 基幹ERP
    "scheduler": "FFEDD5",   # 生産計画・スケジューラ
    "mes":       "EDE9FE",   # 製造実行（MES）
    "plm":       "86EFAC",   # 仕様書・BOM（PLM）
    "portal":    "FCE7F3",   # 取引先向けポータル
    "factory":   "BEF264",   # 工場業務・その他
    "manual":    "F1F5F9",   # 手作業・Excel・紙
    "inhouse":   "CCFBF1",   # 社内DB（FileMaker 等）
}
SYSTEM_LABELS = {
    "sales": "販売・案件管理", "bi": "分析・BI", "erp": "基幹ERP", "scheduler": "生産計画",
    "mes": "製造実行", "plm": "仕様書・BOM", "portal": "取引先ポータル",
    "factory": "工場業務", "manual": "手作業・Excel", "inhouse": "社内DB",
}
INK = "334155"; SUB = "64748B"; RED = "DC2626"; GRAY = "94A3B8"
EDGE_STYLE = {   # kind → (色, 太さpt, 破線)
    "flow":  ("475569", 1.0, False),   # 業務・情報の流れ
    "if":    ("2563EB", 1.0, True),    # システム連携（I/F）
    "goods": ("A16207", 2.0, False),   # 物の流れ
    "paper": ("A16207", 1.5, "dot"),   # 紙の受け渡し
}
NODE_DEFAULT_SIZE = {        # type → (w, h) グリッド単位
    "process": (4, 1.3), "data": (3.4, 2), "doc": (3.4, 1.8),
    "decision": (3.4, 2), "terminal": (3, 1.2), "note": (4, 1.5),
}
SHAPE_OF = {
    "process": MSO_SHAPE.ROUNDED_RECTANGLE, "data": MSO_SHAPE.CAN,
    "doc": MSO_SHAPE.FLOWCHART_DOCUMENT, "decision": MSO_SHAPE.DIAMOND,
    "terminal": MSO_SHAPE.FLOWCHART_TERMINATOR, "note": MSO_SHAPE.RECTANGLE,
}
PAGES = {"A3": (16.54, 11.69), "A4": (11.69, 8.27), "16:9": (13.333, 7.5)}
CXN_IDX = {"T": 0, "L": 1, "B": 2, "R": 3}   # 上記図形はすべて t,l,b,r の4接続点


# ------------------------------------------------------------------ 補助
def _rgb(h): return RGBColor.from_string(h)


def _style_run(run, size, color=INK, bold=False):
    f = run.font
    f.size = Pt(size); f.bold = bold; f.color.rgb = _rgb(color); f.name = FONT
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT)


def _set_text(shape, text, size=10, color=INK, bold=False, align=PP_ALIGN.CENTER,
              anchor=MSO_ANCHOR.MIDDLE, margin=0.04):
    tf = shape.text_frame
    tf.clear(); tf.word_wrap = True; tf.auto_size = None
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = line
        _style_run(r, size, color, bold)


def _no_shadow(shape):
    st = shape._element.find(qn("p:style"))      # テーマ既定の影・線を外す
    if st is not None:
        shape._element.remove(st)


def _line(shape, color=None, width=0.75, dash=False):
    if color is None:
        shape.line.fill.background(); return
    shape.line.color.rgb = _rgb(color); shape.line.width = Pt(width)
    if dash: shape.line.dash_style = MSO_LINE.DASH


def _fill(shape, color):
    if color is None: shape.fill.background()
    else:
        shape.fill.solid(); shape.fill.fore_color.rgb = _rgb(color)


class _Grid:
    def __init__(self, spec):
        pw, ph = PAGES.get(spec.get("page", "A3"), PAGES["A3"])
        g = spec.get("grid", {})
        self.cols, self.rows = g.get("cols", 48), g.get("rows", 30)
        self.left, self.top = 0.4, 1.0
        self.cw = (pw - 0.8) / self.cols
        self.rh = (ph - self.top - 0.4) / self.rows
        self.pw, self.ph = pw, ph

    def box(self, col, row, w, h):   # → EMU (x, y, w, h)
        return (Inches(self.left + col * self.cw), Inches(self.top + row * self.rh),
                Inches(w * self.cw), Inches(h * self.rh))


# ------------------------------------------------------------------ コネクタ
def _port(box, side):
    x, y, w, h = box
    return {"T": (x + w // 2, y), "B": (x + w // 2, y + h),
            "L": (x, y + h // 2), "R": (x + w, y + h // 2)}[side]


def _auto_sides(a, b):
    ax, ay, aw, ah = a; bx, by, bw, bh = b
    acx, acy, bcx, bcy = ax + aw / 2, ay + ah / 2, bx + bw / 2, by + bh / 2
    # 横方向に重なりがなく左右に離れている → 左右接続
    if bx >= ax + aw or ax >= bx + bw:
        return ("R", "L") if bcx > acx else ("L", "R")
    return ("B", "T") if bcy > acy else ("T", "B")


def _route(p0, s0, p1, s1, off):
    """折れ線経路（検査・ラベル位置用）・描画方式・回転・中間線位置(adj)を返す"""
    (x0, y0), (x1, y1) = p0, p1
    h0, h1 = s0 in "LR", s1 in "LR"
    if s0 == s1:                       # 同じ辺どうし → コの字
        if h0:
            mx = max(x0, x1) + off if s0 == "R" else min(x0, x1) - off
            if x1 == x0: x1 += 1
            return [p0, (mx, y0), (mx, y1), p1], "bentConnector3", False, (mx - x0) / (x1 - x0)
        my = max(y0, y1) + off if s0 == "B" else min(y0, y1) - off
        if y1 == y0: y1 += 1
        return [p0, (x0, my), (x1, my), p1], "bentConnector3", True, (my - y0) / (y1 - y0)
    if x0 == x1 or y0 == y1:
        return [p0, p1], "straight", False, None
    if h0 and h1:
        mx = (x0 + x1) // 2
        return [p0, (mx, y0), (mx, y1), p1], "bentConnector3", False, None
    if not h0 and not h1:
        my = (y0 + y1) // 2
        return [p0, (x0, my), (x1, my), p1], "bentConnector3", True, None
    if h0 and not h1:
        return [p0, (x1, y0), p1], "bentConnector2", False, None
    return [p0, (x0, y1), p1], "bentConnector2", True, None


def _set_xfrm(cxn, p0, p1, rotate):
    (x0, y0), (x1, y1) = p0, p1
    Dx, Dy = x1 - x0, y1 - y0
    cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
    if rotate:      # 90°回転：先に縦へ進む経路にする
        dx, dy = Dy, -Dx
    else:
        dx, dy = Dx, Dy
    w, h = abs(dx), abs(dy)
    xfrm = cxn._element.spPr.find(qn("a:xfrm"))
    for k in ("flipH", "flipV", "rot"):
        if k in xfrm.attrib: del xfrm.attrib[k]
    if rotate: xfrm.set("rot", "5400000")
    if dx < 0: xfrm.set("flipH", "1")
    if dy < 0: xfrm.set("flipV", "1")
    xfrm.find(qn("a:off")).set("x", str(int(cx - w / 2)))
    xfrm.find(qn("a:off")).set("y", str(int(cy - h / 2)))
    xfrm.find(qn("a:ext")).set("cx", str(int(w)))
    xfrm.find(qn("a:ext")).set("cy", str(int(h)))


def _connect(cxn, start_shape, s_idx, end_shape, e_idx):
    nv = cxn._element.find(qn("p:nvCxnSpPr")).find(qn("p:cNvCxnSpPr"))
    for tag, shp, idx in (("a:stCxn", start_shape, s_idx), ("a:endCxn", end_shape, e_idx)):
        el = etree.SubElement(nv, qn(tag))
        el.set("id", str(shp.shape_id)); el.set("idx", str(idx))


def _arrow(cxn, color, width, dash):
    _no_shadow(cxn)
    ln = cxn._element.spPr.find(qn("a:ln"))
    if ln is None: ln = etree.SubElement(cxn._element.spPr, qn("a:ln"))
    ln.set("w", str(Pt(width)))
    for c in list(ln): ln.remove(c)
    sf = etree.SubElement(ln, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr")).set("val", color)
    if dash: etree.SubElement(ln, qn("a:prstDash")).set("val", "sysDot" if dash == "dot" else "dash")
    etree.SubElement(ln, qn("a:round"))
    t = etree.SubElement(ln, qn("a:tailEnd")); t.set("type", "triangle")
    t.set("w", "med"); t.set("len", "med")


# ------------------------------------------------------------------ 検査
def _overlap(a, b, pad=0):
    ax, ay, aw, ah = a; bx, by, bw, bh = b
    return ax < bx + bw - pad and bx < ax + aw - pad and ay < by + bh - pad and by < ay + ah - pad


def _seg_hits_box(p, q, box, shrink):
    x, y, w, h = box
    x0, y0, x1, y1 = x + shrink, y + shrink, x + w - shrink, y + h - shrink
    (px, py), (qx, qy) = p, q
    if px == qx:
        return x0 < px < x1 and min(py, qy) < y1 and max(py, qy) > y0
    if py == qy:
        return y0 < py < y1 and min(px, qx) < x1 and max(px, qx) > x0
    return False


# ------------------------------------------------------------------ 本体
def render(spec, out_path):
    if isinstance(spec, str):
        with open(spec, encoding="utf-8") as f: spec = json.load(f)
    spec = deepcopy(spec)
    warn = []
    G = _Grid(spec)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(G.pw), Inches(G.ph)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    meta = spec.get("meta", {})

    # タイトル
    tb = slide.shapes.add_textbox(Inches(0.4), Inches(0.25), Inches(G.pw - 7.8), Inches(0.45))
    _set_text(tb, "■ " + meta.get("title", "業務フロー"), 20, "1E293B", True, PP_ALIGN.LEFT, margin=0)
    info = "　｜　".join(x for x in [
        meta.get("layer", ""), meta.get("state", ""), meta.get("id", ""),
        ("v" + str(meta["version"])) if meta.get("version") else "",
        meta.get("author", ""), meta.get("date", "")] if x)
    tb2 = slide.shapes.add_textbox(Inches(0.4), Inches(0.68), Inches(G.pw - 7.8), Inches(0.25))
    _set_text(tb2, info, 9, SUB, False, PP_ALIGN.LEFT, margin=0)

    # 組織枠
    for o in spec.get("orgs", []):
        x, y, w, h = G.box(o["col"], o.get("row", 0), o["w"], o.get("h", G.rows - o.get("row", 0)))
        s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
        _fill(s, None); _line(s, "94A3B8", 1.0, True); _no_shadow(s)
        if "row" in o and w > h * 2:      # 横長＝スイムレーン：名前は左端に縦書き
            lb = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.32), h)
            _fill(lb, "F1F5F9"); _line(lb, None); _no_shadow(lb)
            _set_text(lb, "\n".join(o["name"]), 12, "1E293B", True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 0.02)
        else:
            _set_text(s, o["name"], 13, "1E293B", False, PP_ALIGN.CENTER, MSO_ANCHOR.TOP, 0.06)

    # システム領域
    sys_boxes = {}
    for sd in spec.get("systems", []):
        box = G.box(sd["col"], sd["row"], sd["w"], sd["h"])
        cat = sd.get("cat", "manual")
        if cat not in SYSTEM_COLORS:
            warn.append(f"[区分] システム {sd.get('id')} の区分 '{cat}' は未定義 → manual で描画"); cat = "manual"
        s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, *box)
        s.adjustments[0] = 0.03
        _fill(s, SYSTEM_COLORS[cat]); _line(s, None); _no_shadow(s)
        _set_text(s, sd.get("name", ""), 11, "1E293B", True, PP_ALIGN.LEFT, MSO_ANCHOR.TOP, 0.08)
        sys_boxes.setdefault(sd["id"], []).append(box)

    # ノード
    nodes, node_boxes = {}, {}
    for n in spec.get("nodes", []):
        t = n.get("type", "process")
        if t not in SHAPE_OF:
            warn.append(f"[種別] ノード {n['id']} の type '{t}' は未定義 → process"); t = "process"
        dw, dh = NODE_DEFAULT_SIZE[t]
        box = G.box(n["col"], n["row"], n.get("w", dw), n.get("h", dh))
        if n["id"] in nodes:
            warn.append(f"[ID] ノードID {n['id']} が重複"); continue
        s = slide.shapes.add_shape(SHAPE_OF[t], *box)
        if t == "process": s.adjustments[0] = 0.18
        if t == "data": s.adjustments[0] = 0.18
        status = n.get("status", "")            # new / changed / removed / ""
        stroke, fcolor, dash, lw = INK if t != "data" else SUB, INK, False, 0.75
        if status in ("new", "changed"): stroke, fcolor, lw = RED, RED, 1.5
        if status == "removed": stroke, fcolor, dash = GRAY, GRAY, True
        _fill(s, "FFFFFF" if t != "terminal" else "1E293B")
        if t == "terminal": fcolor = "FFFFFF"
        if t == "note": _fill(s, "FFFBEB")
        _line(s, stroke, lw, dash); _no_shadow(s)
        size = n.get("size", 13 if n.get("emphasis") else 10)
        _set_text(s, n["label"], size, fcolor, bool(n.get("emphasis")) or t == "terminal")
        nodes[n["id"]] = s; node_boxes[n["id"]] = box
        sysid = n.get("sys")
        if sysid:
            if sysid not in sys_boxes:
                warn.append(f"[ID] ノード {n['id']} の sys '{sysid}' が systems に無い")
            elif not any(_contains(sb, box) for sb in sys_boxes[sysid]):
                warn.append(f"[配置] ノード {n['id']} がシステム領域 {sysid} からはみ出している")

    # ノード同士の重なり
    ids = list(node_boxes)
    for i, a in enumerate(ids):
        for b in ids[i + 1:]:
            if _overlap(node_boxes[a], node_boxes[b]):
                warn.append(f"[重なり] ノード {a} と {b} が重なっている")

    # 矢印
    first_node_el = nodes[ids[0]]._element if ids else None
    used = set(); labels = []
    for k, e in enumerate(spec.get("edges", [])):
        a, b = e.get("from"), e.get("to")
        if a not in nodes or b not in nodes:
            warn.append(f"[ID] 矢印 {a}→{b} の端点が存在しない"); continue
        used.update([a, b])
        s0, s1 = _auto_sides(node_boxes[a], node_boxes[b])
        s0, s1 = e.get("from_side", s0), e.get("to_side", s1)
        p0, p1 = _port(node_boxes[a], s0), _port(node_boxes[b], s1)
        path, prst, rotate, adj = _route(p0, s0, p1, s1, Inches(0.5 * G.cw))
        cxn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT if prst == "straight"
                                         else MSO_CONNECTOR.ELBOW, 0, 0, 1, 1)
        if prst != "straight":
            pg = cxn._element.spPr.find(qn("a:prstGeom")); pg.set("prst", prst)
            if adj is not None:
                av = pg.find(qn("a:avLst"))
                if av is None: av = etree.SubElement(pg, qn("a:avLst"))
                for c in list(av): av.remove(c)
                gd = etree.SubElement(av, qn("a:gd")); gd.set("name", "adj1"); gd.set("fmla", f"val {int(adj * 100000)}")
        x0e, y0e = p0; x1e, y1e = p1
        if adj is not None and s0 in "LR" and x1e == x0e: x1e += 1
        if adj is not None and s0 in "TB" and y1e == y0e: y1e += 1
        _set_xfrm(cxn, p0, (x1e, y1e), rotate)
        _connect(cxn, nodes[a], CXN_IDX[s0], nodes[b], CXN_IDX[s1])
        color, width, dash = EDGE_STYLE.get(e.get("kind", "flow"), EDGE_STYLE["flow"])
        if e.get("status") in ("new", "changed"): color = RED
        _arrow(cxn, color, width, dash)
        # 矢印を図形の背面へ
        if first_node_el is not None:
            first_node_el.addprevious(cxn._element)
        # 横切り検査
        shrink = Inches(0.03)
        for nid, nb in node_boxes.items():
            if nid in (a, b): continue
            if any(_seg_hits_box(path[i], path[i + 1], nb, shrink) for i in range(len(path) - 1)):
                warn.append(f"[横切り] 矢印 {a}→{b} がノード {nid} の上を通過（座標か from_side/to_side を調整）")
        if e.get("label"):
            labels.append((e["label"], path, e.get("label_dx", 0), e.get("label_dy", 0)))

    # 矢印ラベル（最前面）
    for text, path, ldx, ldy in labels:
        segs = [(path[i], path[i + 1]) for i in range(len(path) - 1)]
        p, q = max(segs, key=lambda s: abs(s[0][0] - s[1][0]) + abs(s[0][1] - s[1][1]))
        mx, my = (p[0] + q[0]) / 2, (p[1] + q[1]) / 2
        w = Pt(9 * 1.05 * max(len(l) for l in text.split("\n")) + 8); h = Pt(14 * len(text.split("\n")))
        tb = slide.shapes.add_textbox(int(mx - w / 2 + Inches(ldx * G.cw)), int(my - h / 2 + Inches(ldy * G.rh)), int(w), int(h))
        _fill(tb, "FFFFFF")
        _set_text(tb, text, 8.5, "475569", False, margin=0.01)

    # 孤立ノード
    for nid in ids:
        nd = next(n for n in spec["nodes"] if n["id"] == nid)
        if nid not in used and nd.get("type") != "note" and not nd.get("standalone"):
            warn.append(f"[孤立] ノード {nid} に矢印が1本も無い")

    _legend(slide, G, spec)
    prs.save(out_path)
    return warn


def _contains(outer, inner):
    ox, oy, ow, oh = outer; ix, iy, iw, ih = inner
    return ox <= ix and oy <= iy and ix + iw <= ox + ow and iy + ih <= oy + oh


def _legend(slide, G, spec):
    if spec.get("legend") is False: return
    cats = []
    for sd in spec.get("systems", []):
        c = sd.get("cat", "manual")
        if c not in cats and c in SYSTEM_COLORS: cats.append(c)
    kinds = sorted({e.get("kind", "flow") for e in spec.get("edges", [])}, key=list(EDGE_STYLE).index)
    has_new = any(n.get("status") in ("new", "changed") for n in spec.get("nodes", []))
    items = [("sys", c) for c in cats] + [("edge", k) for k in kinds] + ([("new", None)] if has_new else [])
    if not items: return
    per_row, iw, ih = 5, 1.4, 0.24          # 凡例はヘッダー右側（図の領域外）に置く
    x0 = G.pw - 0.4 - per_row * iw; y0 = 0.2
    names = {"flow": "業務・情報の流れ", "if": "システム連携", "goods": "物の流れ", "paper": "紙の受け渡し"}
    for i, (kind, v) in enumerate(items):
        x = Inches(x0 + (i % per_row) * iw); y = Inches(y0 + (i // per_row) * ih)
        if kind == "sys":
            s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y + Inches(0.04), Inches(0.28), Inches(0.16))
            _fill(s, SYSTEM_COLORS[v]); _line(s, None); _no_shadow(s); label = SYSTEM_LABELS[v]
        elif kind == "edge":
            c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x, y + Inches(0.12), x + Inches(0.28), y + Inches(0.12))
            col, w, d = EDGE_STYLE[v]; _arrow(c, col, w, d); label = names[v]
        else:
            s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y + Inches(0.04), Inches(0.28), Inches(0.16))
            _fill(s, "FFFFFF"); _line(s, RED, 1.5); _no_shadow(s); label = "新規・変更"
        t = slide.shapes.add_textbox(x + Inches(0.32), y, Inches(iw - 0.34), Inches(ih))
        _set_text(t, label, 8, SUB, False, PP_ALIGN.LEFT, margin=0)


if __name__ == "__main__":
    import sys
    ws = render(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else "flow.pptx")
    print("\n".join(ws) if ws else "警告なし")
